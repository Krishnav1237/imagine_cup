"""
PowerPoint processing service.
PURE extraction only.
NO cleaning. NO heuristics.
"""

import io
import logging
from typing import Optional
from pptx import Presentation

logger = logging.getLogger("PPTXProcessor")


class PPTXProcessor:
    """
    Raw PPTX text extractor.
    """

    def extract_text(self, file_data: bytes) -> Optional[str]:
        logger.info("📊 Extracting raw PPTX text")

        try:
            prs = Presentation(io.BytesIO(file_data))
            parts = []

            for slide_idx, slide in enumerate(prs.slides, start=1):
                slide_text = self._extract_slide_text(slide)
                if slide_text:
                    parts.append(f"\n--- Slide {slide_idx} ---\n{slide_text}")

            text = "\n".join(parts)
            logger.info(f"✅ Raw PPTX text length: {len(text)}")
            return text

        except Exception:
            logger.exception("❌ PPTX extraction failed")
            return None

    def _extract_slide_text(self, slide) -> str:
        parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)

            # Tables
            if shape.shape_type == 19:  # TABLE
                parts.append(self._extract_table_text(shape.table))

        return "\n".join(p.strip() for p in parts if p and p.strip())

    def _extract_table_text(self, table) -> str:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        return "\n".join(rows)
